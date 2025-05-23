from abc import ABC, abstractmethod
import docx, openpyxl, csv, PyPDF2
from PIL import Image
from pytesseract import image_to_string

class BaseTextExtractor(ABC):
    @abstractmethod
    def extract_text(self, file):
        pass

class TxtTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        return file.read().decode("utf-8")

class DocxTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        doc = docx.Document(file)
        return "\n".join(para.text for para in doc.paragraphs)

class XlsxTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        workbook = openpyxl.load_workbook(file)
        sheet = workbook.active
        return "\n".join(" ".join(str(cell) for cell in row) for row in sheet.iter_rows(values_only=True))

class CsvTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        reader = csv.reader(file.read().decode('utf-8').splitlines())
        return "\n".join(" ".join(row) for row in reader)

class PdfTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        reader = PyPDF2.PdfReader(file)
        return "\n".join(page.extract_text() for page in reader.pages)

class ImageTextExtractor(BaseTextExtractor):
    def extract_text(self, file):
        image = Image.open(file)
        return image_to_string(image)

class TextExtractorFactory:
    @staticmethod
    def get_extractor(extension):
        extractors = {
            'txt': TxtTextExtractor(),
            'docx': DocxTextExtractor(),
            'xlsx': XlsxTextExtractor(),
            'csv': CsvTextExtractor(),
            'pdf': PdfTextExtractor(),
            'jpg': ImageTextExtractor(),
            'jpeg': ImageTextExtractor(),
            'png': ImageTextExtractor(),
        }
        return extractors.get(extension)